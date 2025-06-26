import csv
import os
import sys
import tempfile
import time
from tempfile import NamedTemporaryFile
from urllib.parse import urlparse

import grpc
import requests
from fastapi import BackgroundTasks, FastAPI, HTTPException
from google.protobuf.json_format import MessageToDict
from openpyxl import load_workbook
from pydantic import BaseModel, Field

# 将 generated 目录及其子目录加入模块搜索路径，以便导入自动生成的 stub
# 假设脚本位于项目根目录，如果不是，请适当调整 ROOT_DIR
ROOT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.append(os.path.join(ROOT_DIR, "generated"))
sys.path.append(os.path.join(ROOT_DIR, "generated", "adminconf", "v1"))

# 导入通过 protoc 生成的 gRPC stub
from knowledgebase_pb2 import FileParseFinishRequest
from knowledgebase_pb2_grpc import KnowledgeBaseServiceStub

# FastAPI app
app = FastAPI()

# 默认配置，可通过环境变量覆盖
GRPC_SERVER_ADDRESS = os.getenv("GRPC_SERVER_ADDRESS", "34.126.174.173:8081")
MINERU_URL = os.getenv("MINERU_URL", "http://localhost:8000/file_parse")
MINERU_OUTPUT_DIR = os.getenv("MINERU_OUTPUT_DIR", "/app/output")


class FileParseRequest(BaseModel):
    file_id: str
    file_url: str = Field(None, description="可选：预签名 URL")
    input_path: str = Field(None, description="可选：本地文件或目录路径，优先使用")
    output_path: str = Field(None, description="可选：解析结果保存路径，文件或目录")


# --- 后台任务 ---
async def run_file_parse_and_callback(req: FileParseRequest):
    """
    这个函数包含了所有的耗时操作，它会在后台运行。
    """
    print(
        "DEBUG: 当前北京时间：",
        time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(time.time() + 8 * 3600)),
    )
    overall_start = time.time()
    print(
        f"DEBUG: Received file_id={req.file_id}, file_url={req.file_url}, input_path={req.input_path}"
    )
    try:
        # 1. 确定来源与下载
        download_start = time.time()
        tmp_path = None  # 初始化 tmp_path
        if req.input_path:
            if not os.path.exists(req.input_path):
                # 后台任务中不能抛出 HTTPException，只能记录日志
                print(f"ERROR: Local path not found: {req.input_path}")
                return  # 提前退出任务
            if os.path.isdir(req.input_path):
                # 批量处理目录
                results = {}
                for fname in sorted(os.listdir(req.input_path)):
                    full = os.path.join(req.input_path, fname)
                    if os.path.isfile(full):
                        print(f"DEBUG: Batch parsing {full}")
                        try:
                            results[fname] = _parse_by_suffix(full)
                        except Exception as e:
                            print(f"ERROR: Failed to parse {fname}: {e}")
                            results[fname] = f"Error parsing file: {e}"

                if req.output_path:
                    os.makedirs(req.output_path, exist_ok=True)
                    for fname, txt in results.items():
                        out = os.path.join(
                            req.output_path, f"{os.path.splitext(fname)[0]}.txt"
                        )
                        print(f"DEBUG: Writing batch output to {out}")
                        with open(out, "w", encoding="utf-8") as f:
                            f.write(txt)
                combined = "\n\n".join(f"=== {k} ===\n{v}" for k, v in results.items())
                download_end = time.time()
                print(
                    f"DEBUG: Download and batch parse time: {download_end - download_start:.2f}s"
                )
                text = combined
                _grpc_finish(req.file_id, text, overall_start, success=True)
                return
            else:
                tmp_path = req.input_path
        elif req.file_url:
            print(f"DEBUG: Downloading from URL {req.file_url}")
            resp = requests.get(req.file_url, stream=True)
            resp.raise_for_status()
            path = urlparse(req.file_url).path
            suffix = os.path.splitext(path)[1].lower()
            with NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                for chunk in resp.iter_content(8192):
                    tmp.write(chunk)
                tmp_path = tmp.name
            print(f"DEBUG: Downloaded to {tmp_path}")
        else:
            print("ERROR: Must provide file_url or input_path")
            _grpc_finish(
                req.file_id,
                "Must provide file_url or input_path",
                overall_start,
                success=False,
            )
            return

        download_end = time.time()
        print(f"DEBUG: Download time: {download_end - download_start:.2f}s")

        # 2. 解析单文件
        parse_start = time.time()
        print(f"DEBUG: Parsing file {tmp_path}")
        text = _parse_by_suffix(tmp_path)
        parse_end = time.time()
        print(f"DEBUG: Parse time: {parse_end - parse_start:.2f}s, length={len(text)}")

        # 3. 写入输出
        write_start = time.time()
        if req.output_path:
            # 确定输出文件名
            if os.path.isdir(req.output_path) or req.output_path.endswith(os.sep):
                # 如果 tmp_path 存在且非空
                base_name = os.path.basename(tmp_path) if tmp_path else "output.txt"
                fn = os.path.splitext(base_name)[0] + ".txt"
                out_file = os.path.join(req.output_path, fn)
            else:
                out_file = req.output_path

            os.makedirs(os.path.dirname(out_file), exist_ok=True)
            print(f"DEBUG: Writing output to {out_file}")
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(text)
        write_end = time.time()
        print(f"DEBUG: Write time: {write_end - write_start:.2f}s")

        # 4. 清理临时文件
        if req.file_url and tmp_path and os.path.exists(tmp_path):
            print(f"DEBUG: Removing temp file {tmp_path}")
            os.remove(tmp_path)

        # 5. gRPC 回调
        grpc_start = time.time()
        print(f"DEBUG: Trigger gRPC FileParseFinish for {req.file_id}")
        _grpc_finish(req.file_id, text, overall_start, success=True)
        grpc_end = time.time()
        print(f"DEBUG: gRPC call time: {grpc_end - grpc_start:.2f}s")
        overall_end = time.time()
        print(f"DEBUG: Total background task time: {overall_end - overall_start:.2f}s")

    except Exception as e:
        # 捕获所有异常，记录日志，并尝试进行失败的 gRPC 回调
        error_message = f"An unexpected error occurred during file processing for file_id {req.file_id}: {e}"
        print(f"ERROR: {error_message}")
        import traceback

        traceback.print_exc()
        # 即使失败，也调用 gRPC 通知对方，并传递错误信息
        _grpc_finish(req.file_id, error_message, overall_start, success=False)
        # 清理可能残留的临时文件
        if (
            req.file_url
            and "tmp_path" in locals()
            and tmp_path
            and os.path.exists(tmp_path)
        ):
            print(f"DEBUG: Removing temp file {tmp_path} after error")
            os.remove(tmp_path)


@app.post("/file_parse_server")
async def file_parse_server(req: FileParseRequest, background_tasks: BackgroundTasks):
    # 将实际的处理函数 `run_file_parse_and_callback` 添加到后台任务队列
    # 先去立即返回成功响应，然后后台处理
    background_tasks.add_task(run_file_parse_and_callback, req)
    print(f"DEBUG: File parse task submitted for {req.file_id}")
    # 立即返回，不等待后台任务完成
    return {"message": "文件解析任务已提交，正在后台处理中"}


def _parse_by_suffix(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        return _parse_txt(path)
    if ext == ".csv":
        return _parse_csv(path)
    if ext in (".xls", ".xlsx"):
        return _parse_excel(path)
    if ext in (".pdf", ".doc", ".docx", ".ppt", ".pptx", ".png", ".jpg", ".jpeg"):
        return _parse_with_mineru(path)
    # 在后台任务中，我们不再抛出 HTTPException
    raise ValueError(f"Unsupported file type: {ext}")


def _parse_csv(path: str) -> str:
    print(f"DEBUG: CSV parse {path}")
    lines = []
    with open(path, newline="", encoding="utf-8") as f:
        for row in csv.reader(f):
            lines.append(",".join(str(c) for c in row))
    return "\n".join(lines)


def _parse_excel(path: str) -> str:
    print(f"DEBUG: Excel parse {path}")
    wb = load_workbook(path, read_only=True)
    out = []
    for s in wb.sheetnames:
        out.append(f"=== {s} ===")
        for r in wb[s].iter_rows(values_only=True):
            out.append(",".join(str(c) if c is not None else "" for c in r))
    return "\n".join(out)


def _parse_txt(path: str) -> str:
    print(f"DEBUG: TXT parse {path}")
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def _parse_with_mineru(path: str) -> str:
    print(f"DEBUG: MinerU parse {path}")
    with open(path, "rb") as f:
        data = {
            "parse_method": "auto",
            "is_json_md_dump": "true",
            "return_images": "false",
            "return_layout": "false",
            "return_info": "false",
            "return_content_list": "false",
            "output_dir": MINERU_OUTPUT_DIR,
        }
        resp = requests.post(MINERU_URL, data=data, files={"file": f})
    resp.raise_for_status()
    js = resp.json()
    print(f"DEBUG: MinerU keys {list(js.keys())}")
    # 确保即使 'md_content' 是 None 或空字符串，也能尝试 'content_list'
    content = js.get("md_content")
    if not content:
        content = "\n".join(js.get("content_list", []))
    return content


def _grpc_finish(name: str, text: str, start_time: float, success: bool) -> dict:
    """
    修改了 gRPC 函数，使其可以指明任务是否成功。
    """
    print(f"DEBUG: gRPC -> {GRPC_SERVER_ADDRESS} for file {name}, success={success}")
    try:
        chan = grpc.insecure_channel(GRPC_SERVER_ADDRESS)
        stub = KnowledgeBaseServiceStub(chan)
        # 根据 success 标志填充请求
        req = FileParseFinishRequest(name=name, text=text, success=success)
        print(f"DEBUG: gRPC req {req}")
        res = stub.FileParseFinish(req)
        print(f"DEBUG: gRPC res {res}")
        elapsed = time.time() - start_time
        print(f"DEBUG: Elapsed since handler start: {elapsed:.2f}s")
        return MessageToDict(res)
    except Exception as e:
        print(f"ERROR: gRPC call failed for file {name}: {e}")
        return {"error": str(e)}


if __name__ == "__main__":
    import uvicorn

    p = int(os.getenv("PORT", 8011))
    print(f"DEBUG: Starting on 0.0.0.0:{p}")
    uvicorn.run(app, host="0.0.0.0", port=p)
